"""Ajoute trois annexes de preuve à une copie de la soutenance Bloc IV.

Les annexes sont placées après la diapositive Questions : elles permettent de
répondre à une demande de preuve sans surcharger le déroulé principal. Les GIFs
déjà présents dans la présentation source ne sont ni remplacés ni modifiés.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "dist" / "bloc4_soutenance_FALIGANT_proposition_v2.pptx"
DEFAULT_OUTPUT = ROOT / "dist" / "bloc4_soutenance_FALIGANT_proposition_v3_preuves.pptx"
EVIDENCE_DIRECTORY = ROOT / "assets" / "evidence"
BACKGROUND = ROOT / "assets" / "style" / "bloc4_slide_background.png"
RASTER_DIRECTORY = ROOT / "dist" / "presentation_evidence_raster"

# Dimensions 16:9 exactes du modèle PowerPoint source, exprimées en EMU.
W = 12192000
H = 6858000
AQUA = RGBColor(67, 198, 182)
MINT = RGBColor(158, 231, 210)
SEAFOAM = RGBColor(232, 248, 244)
SKY = RGBColor(47, 159, 208)
INK = RGBColor(32, 52, 60)
MUTED = RGBColor(100, 121, 130)
LINE = RGBColor(216, 236, 236)
WHITE = RGBColor(255, 255, 255)
WARM = RGBColor(255, 248, 235)


def set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def set_line(shape, color: RGBColor, width: Pt = Pt(1)) -> None:
    shape.line.color.rgb = color
    shape.line.width = width


def add_textbox(
    slide,
    text: str,
    x,
    y,
    w,
    h,
    size: int = 20,
    color: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "Garamond",
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_background(slide) -> None:
    slide.shapes.add_picture(str(BACKGROUND), 0, 0, W, H)


def add_footer(slide, index: int) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.55),
        Inches(7.08),
        Inches(12.23),
        Inches(0.02),
    )
    set_fill(line, MINT)
    add_textbox(slide, "Bloc IV - Soutenance RNCP | Annexe de preuve", Inches(0.62), Inches(7.17), Inches(5.4), Inches(0.17), 9, MUTED)
    add_textbox(
        slide,
        f"A{index:02d}  |  Marc-Alfred FALIGANT - Ynov M2 Data Engineer - 2026",
        Inches(7.5),
        Inches(7.17),
        Inches(5.2),
        Inches(0.17),
        9,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_note(slide, text: str) -> None:
    notes = slide.notes_slide.notes_text_frame
    if notes is not None:
        notes.text = f"Annexe de preuve - non projetée par défaut.\n\n{text}"


def rasterize_svg(svg_name: str) -> Path:
    RASTER_DIRECTORY.mkdir(parents=True, exist_ok=True)
    source = EVIDENCE_DIRECTORY / svg_name
    output = RASTER_DIRECTORY / f"{source.stem}.png"
    if not source.exists():
        raise FileNotFoundError(f"Preuve visuelle introuvable : {source}")
    document = fitz.open(str(source))
    try:
        page = document.load_page(0)
        page.get_pixmap(matrix=fitz.Matrix(1.0, 1.0), alpha=False).save(str(output))
    finally:
        document.close()
    return output


def add_panel(slide, x, y, w, h, heading: str, body: str, tint: RGBColor = SEAFOAM) -> None:
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    set_fill(panel, tint)
    set_line(panel, LINE)
    add_textbox(slide, heading, x + Inches(0.16), y + Inches(0.12), w - Inches(0.32), Inches(0.27), 15, SKY, bold=True)
    add_textbox(slide, body, x + Inches(0.16), y + Inches(0.51), w - Inches(0.32), h - Inches(0.64), 13, INK)


def add_evidence_slide(
    presentation: Presentation,
    index: int,
    title: str,
    eyebrow: str,
    image: Path,
    panels: list[tuple[str, str, RGBColor]],
    conclusion: str,
    note: str,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_background(slide)

    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.53), Inches(0.06), Inches(0.45))
    set_fill(marker, AQUA)
    add_textbox(slide, eyebrow.upper(), Inches(0.78), Inches(0.45), Inches(4.2), Inches(0.23), 10, AQUA, bold=True)
    add_textbox(slide, title, Inches(0.78), Inches(0.72), Inches(11.7), Inches(0.48), 26, SKY)

    slide.shapes.add_picture(str(image), Inches(0.72), Inches(1.45), Inches(7.55), Inches(4.25))
    image_border = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.72), Inches(1.45), Inches(7.55), Inches(4.25))
    image_border.fill.background()
    set_line(image_border, LINE)

    for panel_index, (heading, body, tint) in enumerate(panels):
        add_panel(slide, Inches(8.65), Inches(1.45 + panel_index * 1.36), Inches(3.92), Inches(1.1), heading, body, tint)

    conclusion_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.92), Inches(5.93), Inches(11.66), Inches(0.73))
    set_fill(conclusion_box, WARM)
    set_line(conclusion_box, LINE)
    add_textbox(slide, conclusion, Inches(1.15), Inches(6.13), Inches(11.15), Inches(0.32), 14, INK, align=PP_ALIGN.CENTER)
    add_footer(slide, index)
    add_note(slide, note)


def build(source: Path, output: Path) -> None:
    if not source.exists():
        raise FileNotFoundError(f"Présentation source introuvable : {source}")
    if not BACKGROUND.exists():
        raise FileNotFoundError(f"Fond de présentation introuvable : {BACKGROUND}")

    open_meteo = rasterize_svg("open_meteo_collection_snapshot.svg")
    github_projects = rasterize_svg("github_projects_snapshot.svg")
    github_actions = rasterize_svg("github_actions_snapshot.svg")

    presentation = Presentation(str(source))
    if (presentation.slide_width, presentation.slide_height) != (W, H):
        raise ValueError("Le format de la présentation source ne correspond pas au format 16:9 attendu.")

    add_evidence_slide(
        presentation,
        1,
        "API externe : le brut est conservé avant toute transformation",
        "Annexe 1 - Bloc I",
        open_meteo,
        [
            ("Collecte HTTP", "Open-Meteo interrogée le 7 septembre 2026.", SEAFOAM),
            ("Brut et intégrité", "JSON horodaté, empreinte SHA-256, aucune réécriture du payload.", RGBColor(239, 248, 253)),
            ("Chargement", "Ligne retrouvée dans raw.open_meteo_snapshot.", WARM),
        ],
        "Elle apporte un contexte météo aux excursions thermiques ; elle ne prend aucune décision métier à la place de la qualité.",
        "À ouvrir si le jury demande comment une source API externe est réellement collectée et tracée. Montrer le JSON brut, l'empreinte et la table raw.",
    )
    add_evidence_slide(
        presentation,
        2,
        "Pilotage : les jalons existent dans un outil de suivi renseigné",
        "Annexe 2 - Bloc III",
        github_projects,
        [
            ("Huit jalons", "De PIL-01 à PIL-08, liés aux blocs et à la soutenance.", SEAFOAM),
            ("Responsabilité", "Un responsable, une échéance et un statut par action.", RGBColor(239, 248, 253)),
            ("État observé", "Sept éléments terminés ; recette finale encore planifiée.", WARM),
        ],
        "Le Gantt donne une trajectoire ; GitHub Projects apporte le suivi vivant des responsabilités, dates et preuves.",
        "À ouvrir si le jury demande comment le projet a été suivi. Insister sur les jalons, les dates et la différence entre terminé et planifié.",
    )
    add_evidence_slide(
        presentation,
        3,
        "CI/CD : une exécution distante produit des artefacts vérifiables",
        "Annexe 3 - Bloc IV",
        github_actions,
        [
            ("Déclenchement", "Run workflow_dispatch exécuté manuellement et terminé avec succès.", SEAFOAM),
            ("Qualité", "10 tests Python, 20 contrôles dbt et artefact dbt conservé.", RGBColor(239, 248, 253)),
            ("Livraison", "Image OCI publiée avec un digest immuable, SBOM et provenance.", WARM),
        ],
        "Le résultat ne dépend pas d'une démonstration locale : le workflow rejoue les contrôles et publie une image identifiée par son digest.",
        "À ouvrir si le jury demande une preuve de CI/CD. Montrer workflow_dispatch, deux jobs réussis, artefact dbt et digest OCI ; distinguer un tag d'un digest immuable.",
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    print(f"Présentation enrichie : {output}")


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ajoute les annexes de preuve à une copie de la soutenance.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_arguments()
    build(arguments.source, arguments.output)
