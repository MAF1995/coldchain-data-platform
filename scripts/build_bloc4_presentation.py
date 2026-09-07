"""Build the editable Bloc IV oral presentation as a Canva-compatible PPTX."""

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist" / "bloc4_soutenance_FALIGANT.pptx"
SCREENSHOTS = ROOT / "assets" / "screenshots" / "bloc4"
LOGOS = ROOT / "assets" / "logos"
BACKGROUND = ROOT / "assets" / "style" / "bloc4_slide_background.png"
DIAGRAMS = ROOT / "assets" / "diagrams" / "rendered"

W, H = Inches(13.333), Inches(7.5)
AQUA = RGBColor(67, 198, 182)
MINT = RGBColor(158, 231, 210)
SEAFOAM = RGBColor(232, 248, 244)
SKY = RGBColor(47, 159, 208)
INK = RGBColor(32, 52, 60)
MUTED = RGBColor(100, 121, 130)
LINE = RGBColor(216, 236, 236)
WHITE = RGBColor(255, 255, 255)
ROAD = RGBColor(224, 231, 233)
WARM = RGBColor(255, 240, 211)


def rgb(color: RGBColor) -> RGBColor:
    return color


def set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)


def set_line(shape, color: RGBColor, width: Pt = Pt(1)) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = width


def ensure_background() -> Path:
    """Create a subtle linear clinical gradient used behind every slide."""
    BACKGROUND.parent.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1600, 900), (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    for x in range(1600):
        progress = x / 1599
        green = min(1.0, progress * 1.15)
        blue = max(0.0, (progress - 0.55) / 0.45)
        red = int(255 - 8 * green - 3 * blue)
        g = int(255 - 2 * green - 4 * blue)
        b = int(255 - 5 * green)
        draw.line((x, 0, x, 900), fill=(red, g, b))
    draw.rectangle((0, 820, 1600, 900), fill=(238, 249, 247))
    canvas.save(BACKGROUND)
    return BACKGROUND


def add_background(slide) -> None:
    slide.shapes.add_picture(str(ensure_background()), 0, 0, W, H)


def add_presenter_note(slide, timing: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.text = f"Repère de présentation : {timing}\n\nVoir la trame orale du bloc IV pour les éléments à commenter."


def add_textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False,
                align=PP_ALIGN.LEFT, font="Garamond", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return box


def add_bullets(slide, items, x, y, w, h, size=19, color=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(5)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Garamond"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = rgb(color)
        paragraph.space_after = Pt(11)
        paragraph.bullet = True
    return box


def add_footer(slide, index: int) -> None:
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(7.08), Inches(12.23), Inches(0.02))
    set_fill(line, MINT)
    add_textbox(slide, "Bloc IV - Soutenance RNCP", Inches(0.62), Inches(7.17), Inches(4.4), Inches(0.17), 9, MUTED)
    add_textbox(slide, f"{index:02d}  |  Marc-Alfred FALIGANT - Ynov M2 Data Engineer - 2026", Inches(8.0), Inches(7.17), Inches(4.7), Inches(0.17), 9, MUTED, align=PP_ALIGN.RIGHT)


def base_slide(prs: Presentation, title: str, index: int, phase: str, kicker: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.53), Inches(0.06), Inches(0.45))
    set_fill(slide.shapes[-1], AQUA)
    if kicker:
        add_textbox(slide, kicker.upper(), Inches(0.78), Inches(0.45), Inches(4), Inches(0.25), 10, AQUA, bold=True)
    add_textbox(slide, title, Inches(0.78), Inches(0.72), Inches(11.7), Inches(0.55), 27, SKY)
    add_footer(slide, index)
    add_presenter_note(slide, phase)
    return slide


def add_panel(slide, x, y, w, h, heading: str, body: str, tint=SEAFOAM, heading_color=SKY, body_size=17):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    set_fill(panel, tint)
    set_line(panel, LINE)
    add_textbox(slide, heading, x + Inches(0.18), y + Inches(0.16), w - Inches(0.36), Inches(0.3), 16, heading_color, bold=True)
    add_textbox(slide, body, x + Inches(0.18), y + Inches(0.58), w - Inches(0.36), h - Inches(0.72), body_size, INK)
    return panel


def add_metric(slide, x, y, value: str, label: str):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(3.58), Inches(1.1))
    set_fill(box, SEAFOAM)
    set_line(box, LINE)
    add_textbox(slide, value, x + Inches(0.2), y + Inches(0.13), Inches(3.1), Inches(0.44), 27, SKY)
    add_textbox(slide, label, x + Inches(0.2), y + Inches(0.64), Inches(3.15), Inches(0.28), 14, INK)


def add_image(slide, path: Path, x, y, w, h):
    slide.shapes.add_picture(str(path), x, y, w, h)
    border = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h)
    border.fill.background()
    set_line(border, LINE)


def add_logo(slide, name: str, x, y, w, h):
    """Place a logo without distorting its aspect ratio."""
    path = LOGOS / f"{name}.png"
    with Image.open(path) as image:
        ratio = image.width / image.height
    if w / h > ratio:
        actual_h = h
        actual_w = h * ratio
    else:
        actual_w = w
        actual_h = w / ratio
    slide.shapes.add_picture(
        str(path),
        x + (w - actual_w) / 2,
        y + (h - actual_h) / 2,
        actual_w,
        actual_h,
    )


def add_tech_node(slide, x, y, role: str, logo: str | None, tint=SEAFOAM):
    node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(2.08), Inches(1.18))
    set_fill(node, tint)
    set_line(node, LINE)
    if logo:
        add_logo(slide, logo, x + Inches(0.2), y + Inches(0.12), Inches(1.68), Inches(0.52))
    else:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Inches(0.82), y + Inches(0.14), Inches(0.45), Inches(0.45))
        set_fill(dot, AQUA)
        add_textbox(slide, "capteur", x + Inches(0.2), y + Inches(0.25), Inches(1.68), Inches(0.2), 10, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, role, x + Inches(0.13), y + Inches(0.75), Inches(1.82), Inches(0.25), 12, INK, align=PP_ALIGN.CENTER)


def add_arrow(slide, x1, y1, x2, y2, color=SKY):
    arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    arrow.line.color.rgb = rgb(color)
    arrow.line.width = Pt(1.8)
    arrow.line.end_arrowhead = True


def add_architecture(slide, y=Inches(2.35)):
    nodes = [
        ("Signal machine", None, WARM),
        ("Entrée terrain", "mqtt", SEAFOAM),
        ("Bus d'événements", "kafka", WARM),
        ("Référentiel de données", "postgres", SEAFOAM),
        ("Pilotage", "grafana", WARM),
    ]
    left = Inches(0.85)
    node_w = Inches(2.1)
    for idx, (label, logo, tint) in enumerate(nodes):
        x = left + idx * Inches(2.48)
        add_tech_node(slide, x, y, label, logo, tint)
        if idx < len(nodes) - 1:
            add_arrow(slide, x + node_w, y + Inches(0.59), x + Inches(2.48), y + Inches(0.59))
    add_textbox(slide, "Temps réel", Inches(3.35), y + Inches(1.38), Inches(2.3), Inches(0.25), 11, MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Qualité et traçabilité", Inches(8.1), y + Inches(1.38), Inches(2.8), Inches(0.25), 11, MUTED, align=PP_ALIGN.CENTER)


def add_service_group(slide, x, y, title: str, items):
    add_textbox(slide, title, x, y, Inches(3.65), Inches(0.28), 17, SKY, bold=True, align=PP_ALIGN.CENTER)
    for index, (logo, role) in enumerate(items):
        tile_x = x + Inches(index * 1.22)
        tile = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, tile_x, y + Inches(0.48), Inches(1.08), Inches(1.5))
        set_fill(tile, SEAFOAM)
        set_line(tile, LINE)
        add_logo(slide, logo, tile_x + Inches(0.13), y + Inches(0.6), Inches(0.82), Inches(0.55))
        add_textbox(slide, role, tile_x + Inches(0.08), y + Inches(1.23), Inches(0.92), Inches(0.2), 9, INK, align=PP_ALIGN.CENTER)


def add_site_context(slide):
    add_textbox(slide, "Une situation de production à périmètre anonymisé", Inches(0.8), Inches(1.4), Inches(7.2), Inches(0.4), 21, INK)
    add_textbox(slide, "Les zones, identifiants machine et lots sont pseudonymisés dans les preuves.", Inches(0.8), Inches(1.95), Inches(6.5), Inches(0.8), 18, MUTED)
    # Illustration deliberately schematic rather than an asserted external site image.
    canvas = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.75), Inches(1.35), Inches(4.5), Inches(4.85))
    set_fill(canvas, RGBColor(246, 249, 249))
    set_line(canvas, LINE)
    for x, y, w, h in [(8.1, 1.8, 1.2, 0.65), (9.65, 1.85, 1.8, 0.75), (8.45, 3.05, 1.7, 0.85), (10.55, 3.2, 1.15, 0.95), (8.25, 4.55, 1.35, 0.7), (10.0, 4.75, 1.55, 0.65)]:
        building = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        set_fill(building, WHITE)
        set_line(building, RGBColor(190, 204, 207))
    road_h = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(7.85), Inches(2.67), Inches(4.3), Inches(0.22))
    set_fill(road_h, ROAD)
    road_v = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(9.65), Inches(1.5), Inches(0.24), Inches(4.4))
    set_fill(road_v, ROAD)
    zone = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.22), Inches(4.46), Inches(3.5), Inches(1.1))
    zone.fill.background()
    set_line(zone, AQUA, Pt(2.2))
    add_textbox(slide, "Périmètre\nanonymisé", Inches(8.55), Inches(4.75), Inches(2.85), Inches(0.45), 16, AQUA, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Schéma de situation à vocation illustrative.", Inches(7.9), Inches(6.35), Inches(4.2), Inches(0.25), 10, MUTED, align=PP_ALIGN.CENTER)


def add_roadmap(slide):
    entries = [
        ("I", "Contrat MQTT", AQUA),
        ("II", "PostgreSQL + dbt", SKY),
        ("III", "Kafka + Airflow", AQUA),
        ("IV", "Spark + supervision", SKY),
        ("V", "Durcissement production", AQUA),
    ]
    y = Inches(2.8)
    x = Inches(0.85)
    for index, (roman, label, color) in enumerate(entries):
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + index * Inches(2.45), y, Inches(0.64), Inches(0.64))
        set_fill(circle, color)
        add_textbox(slide, roman, x + index * Inches(2.45), y + Inches(0.15), Inches(0.64), Inches(0.25), 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
        if index < len(entries) - 1:
            connector = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + index * Inches(2.45) + Inches(0.64), y + Inches(0.3), Inches(1.8), Inches(0.05))
            set_fill(connector, MINT)
        add_textbox(slide, label, x + index * Inches(2.45) - Inches(0.25), y + Inches(0.85), Inches(1.15), Inches(0.55), 13, INK, align=PP_ALIGN.CENTER)
    add_panel(slide, Inches(1.1), Inches(4.55), Inches(11.1), Inches(1.35), "Décision", "Le socle est démontré. L'étape suivante n'est pas d'ajouter des outils, mais de durcir les identités, la haute disponibilité, les sauvegardes et la recette de préproduction.", SEAFOAM, body_size=14)


def build() -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    for layout in prs.slide_layouts:
        layout.name = layout.name

    # 1. Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), W, Inches(0.18))
    set_fill(band, AQUA)
    add_textbox(slide, "BLOC IV", Inches(0.8), Inches(1.25), Inches(2), Inches(0.35), 14, AQUA, bold=True)
    add_textbox(slide, "Concevoir et opérer\nune infrastructure data", Inches(0.8), Inches(1.75), Inches(7.0), Inches(1.25), 34, SKY)
    add_textbox(slide, "Plateforme de supervision de la chaîne du froid pharmaceutique", Inches(0.8), Inches(3.25), Inches(7.0), Inches(0.5), 22, INK)
    add_textbox(slide, "Marc-Alfred FALIGANT\nYnov - M2 Data Engineer - 2026", Inches(0.8), Inches(4.35), Inches(4.0), Inches(0.65), 17, MUTED)
    cover_shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(1.55), Inches(3.6), Inches(3.6))
    set_fill(cover_shape, SEAFOAM)
    set_line(cover_shape, LINE)
    add_textbox(slide, "PÉRIMÈTRE", Inches(8.9), Inches(2.1), Inches(2.7), Inches(0.35), 18, AQUA, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Périmètre industriel\nanonymisé", Inches(8.9), Inches(2.75), Inches(2.7), Inches(0.6), 24, SKY, align=PP_ALIGN.CENTER)
    add_textbox(slide, "Signal machine\n→ donnée traçable\n→ indicateur qualité", Inches(8.9), Inches(3.75), Inches(2.7), Inches(0.75), 17, INK, align=PP_ALIGN.CENTER)
    add_footer(slide, 1)
    add_presenter_note(slide, "00:00 - 00:30")

    # 2. Context
    slide = base_slide(prs, "Le besoin : rendre la chaîne du froid explicable", 2, "00:00 - 02:00", "Mise en situation")
    add_panel(slide, Inches(0.85), Inches(1.7), Inches(3.7), Inches(3.6), "Périmètre", "Zone industrielle anonymisée. Les identifiants opérationnels, les lots et les emplacements sont pseudonymisés dans les preuves.", SEAFOAM, body_size=19)
    add_panel(slide, Inches(4.8), Inches(1.7), Inches(3.7), Inches(3.6), "Événement", "Les équipements remontent des mesures de température et des signaux de fonctionnement. La valeur du projet est de conserver le contexte du message, pas seulement sa valeur numérique.", RGBColor(239, 248, 253), body_size=19)
    add_panel(slide, Inches(8.75), Inches(1.7), Inches(3.7), Inches(3.6), "Décision attendue", "Identifier une excursion, vérifier la fraîcheur du flux, comprendre son origine et disposer d'un mécanisme de reprise lorsqu'un composant échoue.", RGBColor(255, 248, 235), body_size=19)
    add_textbox(slide, "Les éléments de contexte et les identifiants opérationnels sont anonymisés. Les résultats présentés proviennent de scénarios contrôlés de validation.", Inches(1.05), Inches(5.95), Inches(11.1), Inches(0.35), 14, MUTED, align=PP_ALIGN.CENTER)

    # 3. Problem and goals
    slide = base_slide(prs, "Une donnée machine doit devenir une preuve qualité", 3, "02:00 - 04:00", "Enjeu")
    add_panel(slide, Inches(0.8), Inches(1.65), Inches(3.7), Inches(3.75), "Problème", "Des messages machines arrivent en continu. Sans contrat, horodatage, stockage brut et supervision, il est impossible d'expliquer une excursion de température ou de rejouer un incident.", SEAFOAM, body_size=19)
    add_panel(slide, Inches(4.8), Inches(1.65), Inches(3.7), Inches(3.75), "Objectif", "Relier les signaux d'équipement aux indicateurs de conformité, sans perdre la trace du message d'origine ni créer de doublon lors d'un rejeu.", RGBColor(239, 248, 253), body_size=19)
    add_panel(slide, Inches(8.8), Inches(1.65), Inches(3.7), Inches(3.75), "Limite assumée", "La plateforme démontre le fonctionnement et la recette. Elle ne constitue pas une qualification de production ni une décision de libération de lot.", RGBColor(255, 248, 235), body_size=19)

    # 4. Architecture
    slide = base_slide(prs, "L'architecture : découpler, tracer, contrôler", 4, "04:00 - 06:00", "Architecture")
    add_architecture(slide)
    add_panel(slide, Inches(0.95), Inches(4.1), Inches(3.65), Inches(1.1), "Découpler", "Séparer la réception de la consommation.", SEAFOAM, body_size=16)
    add_panel(slide, Inches(4.85), Inches(4.1), Inches(3.65), Inches(1.1), "Tracer", "Conserver le message source et ses métadonnées.", RGBColor(239, 248, 253), body_size=16)
    add_panel(slide, Inches(8.75), Inches(4.1), Inches(3.65), Inches(1.1), "Contrôler", "Rendre visibles qualité, délai et incidents.", RGBColor(255, 248, 235), body_size=16)

    # 5. Containers
    slide = base_slide(prs, "Docker : rendre le socle reproductible", 5, "06:00 - 07:30", "Exploitation")
    add_service_group(slide, Inches(0.85), Inches(1.75), "Flux et persistance", [("mqtt", "entrée"), ("kafka", "rejeu"), ("postgres", "stockage")])
    add_service_group(slide, Inches(4.85), Inches(1.75), "Transformation", [("airflow", "orchestrer"), ("dbt", "contrôler"), ("spark", "calculer")])
    add_service_group(slide, Inches(8.85), Inches(1.75), "Exploitation", [("prometheus", "alerter"), ("grafana", "voir"), ("docker", "isoler")])
    add_architecture(slide, Inches(4.15))
    add_textbox(slide, "Docker rend le socle reproductible : chaque composant est isolé, démarré avec sa configuration et observable séparément.", Inches(1.0), Inches(6.05), Inches(11.2), Inches(0.3), 15, MUTED, align=PP_ALIGN.CENTER)

    # 6. Real-time flow
    slide = base_slide(prs, "Méthode I : le flux temps réel", 6, "07:30 - 09:00", "Streaming")
    add_image(slide, DIAGRAMS / "08_bloc4_realtime_relationships.png", Inches(0.95), Inches(2.0), Inches(11.45), Inches(2.25))
    add_panel(slide, Inches(1.1), Inches(4.65), Inches(3.45), Inches(1.12), "Contrat", "Le message est validé avant de poursuivre.", SEAFOAM, body_size=13)
    add_panel(slide, Inches(4.95), Inches(4.65), Inches(3.45), Inches(1.12), "Reprise", "Écrire puis confirmer l'offset.", RGBColor(239, 248, 253), body_size=13)
    add_panel(slide, Inches(8.8), Inches(4.65), Inches(3.45), Inches(1.12), "Quarantaine", "Un message invalide reste explicable.", RGBColor(255, 248, 235), body_size=13)

    # 7. Grafana
    slide = base_slide(prs, "La preuve du flux : les compteurs se réconcilient", 7, "09:00 - 11:00", "Streaming")
    add_image(slide, SCREENSHOTS / "01_grafana_pipeline_monitoring.png", Inches(0.85), Inches(1.5), Inches(8.0), Inches(4.5))
    add_metric(slide, Inches(9.15), Inches(1.65), "60 = 60 = 60", "réception MQTT, Kafka, PostgreSQL")
    add_metric(slide, Inches(9.15), Inches(3.0), "0", "message en retard à la capture")
    add_metric(slide, Inches(9.15), Inches(4.35), "738 ms", "latence p95 observée")
    add_textbox(slide, "Cette séquence valide la continuité et l'observabilité. Elle ne remplace pas un test de charge industriel.", Inches(0.95), Inches(6.2), Inches(11.9), Inches(0.3), 13, MUTED, align=PP_ALIGN.CENTER)

    # 8. Data quality
    slide = base_slide(prs, "La qualité : contrôler avant de publier", 8, "11:00 - 12:30", "Entrepôt")
    add_image(slide, DIAGRAMS / "09_bloc4_quality_relationships.png", Inches(0.85), Inches(1.65), Inches(11.7), Inches(3.28))
    add_metric(slide, Inches(1.15), Inches(5.95), "527 / 527", "brut réconcilié avec les faits")
    add_metric(slide, Inches(4.88), Inches(5.95), "20 / 20", "tests dbt réussis")
    add_metric(slide, Inches(8.61), Inches(5.95), "3", "excursions qualifiées")

    # 9. Airflow
    slide = base_slide(prs, "Méthode II : Airflow orchestre la recette quotidienne", 9, "12:30 - 14:30", "Orchestration")
    add_image(slide, SCREENSHOTS / "05_airflow_dag_execution.png", Inches(0.85), Inches(1.5), Inches(8.2), Inches(4.62))
    add_bullets(slide, ["Contrôle de la source brute", "Chargement du référentiel", "Construction dbt et tests", "Barrière qualité", "Publication de la preuve"], Inches(9.35), Inches(1.75), Inches(3.0), Inches(3.4), 17)
    add_textbox(slide, "Six tâches réussies, aucun échec sur l'exécution présentée.", Inches(9.2), Inches(5.6), Inches(3.2), Inches(0.35), 15, AQUA, bold=True, align=PP_ALIGN.CENTER)

    # 10. Spark
    slide = base_slide(prs, "Méthode III : Spark prépare les calculs volumineux", 10, "14:30 - 16:30", "Calcul distribué")
    add_image(slide, SCREENSHOTS / "04_spark_cluster_execution.png", Inches(0.85), Inches(1.5), Inches(8.15), Inches(4.57))
    add_metric(slide, Inches(9.25), Inches(1.75), "2", "workers disponibles")
    add_metric(slide, Inches(9.25), Inches(3.1), "1,8 min", "application terminée")
    add_metric(slide, Inches(9.25), Inches(4.45), "6", "agrégats Parquet produits")
    add_textbox(slide, "Spark n'est pas nécessaire pour 527 lignes. Il montre la méthode distribuée et devient pertinent lors des consolidations massives.", Inches(0.95), Inches(6.2), Inches(11.8), Inches(0.3), 13, MUTED, align=PP_ALIGN.CENTER)

    # 11. Prometheus targets
    slide = base_slide(prs, "Superviser : savoir si chaque maillon répond", 11, "16:30 - 18:00", "Observabilité")
    add_image(slide, SCREENSHOTS / "02_prometheus_targets.png", Inches(0.85), Inches(1.5), Inches(8.25), Inches(4.65))
    add_panel(slide, Inches(9.45), Inches(1.75), Inches(2.8), Inches(1.1), "Cible 1", "Bridge MQTT vers Kafka", SEAFOAM, body_size=15)
    add_panel(slide, Inches(9.45), Inches(3.0), Inches(2.8), Inches(1.1), "Cible 2", "Consommateur Kafka vers PostgreSQL", SEAFOAM, body_size=15)
    add_panel(slide, Inches(9.45), Inches(4.25), Inches(2.8), Inches(1.1), "Cible 3", "Prometheus", SEAFOAM, body_size=15)

    # 12. Alerts
    slide = base_slide(prs, "Alerter : définir qui doit agir et pourquoi", 12, "18:00 - 19:30", "Observabilité")
    add_image(slide, SCREENSHOTS / "03_prometheus_alerts.png", Inches(0.85), Inches(1.55), Inches(7.8), Inches(4.4))
    add_bullets(slide, ["Composant indisponible", "Aucun message MQTT récent", "Lag Kafka trop élevé", "Erreur de livraison", "Activité de la DLQ"], Inches(9.0), Inches(1.8), Inches(3.0), Inches(3.5), 17)
    add_textbox(slide, "Les règles sont bien chargées. Elles sont inactives pendant la recette, ce qui est le résultat attendu dans cette situation normale.", Inches(0.95), Inches(6.2), Inches(11.7), Inches(0.3), 13, MUTED, align=PP_ALIGN.CENTER)

    # 13. Security
    slide = base_slide(prs, "Sécurité et traçabilité : les conditions de confiance", 13, "19:30 - 21:00", "Maîtrise")
    add_panel(slide, Inches(0.85), Inches(1.65), Inches(3.65), Inches(3.45), "Intégrité", "Contrat JSON, empreinte raw_hash, idempotence et validation de la structure.", SEAFOAM, body_size=19)
    add_panel(slide, Inches(4.85), Inches(1.65), Inches(3.65), Inches(3.45), "Confidentialité", "Site et zones pseudonymisés, droits séparés par rôle, secrets à externaliser avant production.", RGBColor(239, 248, 253), body_size=19)
    add_panel(slide, Inches(8.85), Inches(1.65), Inches(3.65), Inches(3.45), "Auditabilité", "Brut conservé, transformations dbt documentées, preuve datée, exécutions Airflow et supervision traçables.", RGBColor(255, 248, 235), body_size=19)

    # 14. Incident
    slide = base_slide(prs, "Incident traité : Spark ne pouvait pas écrire ses journaux", 14, "21:00 - 23:00", "Retour d'expérience")
    add_panel(slide, Inches(0.85), Inches(1.7), Inches(3.55), Inches(3.8), "Symptôme", "Le premier spark-submit échoue : le volume d'event logs possède un propriétaire incompatible avec l'utilisateur Spark 185.", RGBColor(255, 248, 235), body_size=19)
    add_panel(slide, Inches(4.85), Inches(1.7), Inches(3.55), Inches(3.8), "Méthode", "Lecture des logs, localisation du refus, contrôle de l'UID, reproduction sur un volume neuf, validation driver/workers.", SEAFOAM, body_size=19)
    add_panel(slide, Inches(8.85), Inches(1.7), Inches(3.55), Inches(3.8), "Correction", "Initialisation du volume avec chown 185:185, puis configuration explicite du driver. Rejeu réussi, sans perte du brut.", RGBColor(239, 248, 253), body_size=19)

    # 15. Verification
    slide = base_slide(prs, "La recette : des contrôles variés, pas une seule commande", 15, "23:00 - 24:30", "Validation")
    checks = [("Python", "9/9", "contrats et normalisation"), ("dbt", "20/20", "qualité et relations"), ("Temps réel", "60=60=60", "réconciliation des compteurs"), ("Airflow", "6 vertes", "orchestration et barrière"), ("Spark", "FINISHED", "calcul distribué"), ("Prometheus", "3 UP", "cibles supervisées")]
    for i, (name, value, detail) in enumerate(checks):
        row, col = divmod(i, 3)
        x, y = Inches(0.9 + col * 4.1), Inches(1.6 + row * 2.05)
        add_panel(slide, x, y, Inches(3.65), Inches(1.55), name, f"{value}\n{detail}", SEAFOAM, body_size=17)

    # 16. Roadmap
    slide = base_slide(prs, "Roadmap : le prochain travail est le durcissement", 16, "24:30 - 26:30", "Trajectoire")
    add_roadmap(slide)

    # 17. Decision
    slide = base_slide(prs, "Ce qui est prêt, ce qui reste à qualifier", 17, "26:30 - 28:30", "Bilan")
    add_panel(slide, Inches(0.85), Inches(1.65), Inches(5.5), Inches(3.95), "Démontré", "Flux temps réel découplé, stockage brut et analytique, dbt, Airflow, Spark, supervision, alertes, recette et incident réellement résolu.", SEAFOAM, body_size=20)
    add_panel(slide, Inches(6.95), Inches(1.65), Inches(5.5), Inches(3.95), "Avant production", "TLS et ACL MQTT/Kafka, haute disponibilité, gestionnaire de secrets, sauvegardes restaurées, test de charge, SLO/RPO/RTO contractualisés et recette de préproduction.", RGBColor(255, 248, 235), body_size=20)

    # 18. Conclusion
    slide = base_slide(prs, "Conclusion", 18, "28:30 - 30:00", "Clôture")
    add_textbox(slide, "Le projet transforme un signal industriel en information vérifiable, rejouable et supervisée.", Inches(1.05), Inches(1.7), Inches(11.2), Inches(0.65), 27, SKY, align=PP_ALIGN.CENTER)
    add_bullets(slide, ["Une architecture adaptée aux rythmes différents : temps réel, orchestration, calcul distribué.", "Une qualité observable : tests, réconciliation, barrières et alertes.", "Une trajectoire réaliste : validation reproductible aujourd'hui, durcissement avant production."], Inches(1.45), Inches(3.0), Inches(10.4), Inches(1.9), 20)
    add_textbox(slide, "Questions", Inches(0.9), Inches(5.8), Inches(11.4), Inches(0.5), 28, AQUA, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"PPTX généré : {OUT}")


if __name__ == "__main__":
    build()
